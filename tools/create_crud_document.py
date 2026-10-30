from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

OUT = "docs/Documentacao_CRUD_Spring_Boot.docx"
NAVY = "16324F"
BLUE = "2D6A9F"
PALE = "EAF2F8"
LIGHT = "F6F8FA"
GRID = "D9E1E8"
TEXT = RGBColor(31, 41, 55)

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)

def set_cell_margins(cell, top=110, start=130, bottom=110, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

def set_cell_border(cell, color=GRID):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        elem = borders.find(tag)
        if elem is None:
            elem = OxmlElement(f"w:{edge}")
            borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), "6")
        elem.set(qn("w:color"), color)

def set_width(cell, cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(cm * 567)))
    tc_w.set(qn("w:type"), "dxa")

def add_text(p, text, bold=False, color=None, size=None):
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Aptos"
    r._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    if size:
        r.font.size = Pt(size)
    return r

def para(doc, text="", style=None, space_after=6):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(space_after)
    add_text(p, text)
    return p

def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    add_text(p, text, color=NAVY, size=15 if level == 1 else 12)
    return p

def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        add_text(p, item)

def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        set_width(cell, widths[i]); set_cell_shading(cell, NAVY); set_cell_margins(cell); set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, h, bold=True, color="FFFFFF", size=8.5)
    for row_num, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            set_width(cell, widths[i]); set_cell_margins(cell); set_cell_border(cell)
            if row_num % 2 == 1: set_cell_shading(cell, LIGHT)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i >= len(row)-2 else WD_ALIGN_PARAGRAPH.LEFT
            add_text(p, value, size=8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(2.0); sec.bottom_margin = Cm(1.8); sec.left_margin = Cm(2.0); sec.right_margin = Cm(2.0)

styles = doc.styles
styles["Normal"].font.name = "Aptos"; styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos"); styles["Normal"].font.size = Pt(10)
styles["Normal"].font.color.rgb = TEXT
for s in ("Title", "Heading 1", "Heading 2"):
    styles[s].font.name = "Aptos Display" if s == "Title" else "Aptos"
    styles[s].font.color.rgb = RGBColor.from_string(NAVY)
styles["Title"].font.size = Pt(28); styles["Heading 1"].font.size = Pt(15); styles["Heading 2"].font.size = Pt(12)

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_text(footer, "Mottainai | Documentação técnica do projeto", color="6B7280", size=8)

# Cover
p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(90); p.paragraph_format.space_after = Pt(10)
add_text(p, "DOCUMENTAÇÃO TÉCNICA", color=BLUE, size=11, bold=True)
p = doc.add_paragraph(style="Title"); p.paragraph_format.space_after = Pt(8)
add_text(p, "CRUD administrativo com Spring Boot", color="000000", size=28)
p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(20)
add_text(p, "Projeto Mottainai", color="4B5563", size=15)

intro = doc.add_table(rows=3, cols=2); intro.alignment = WD_TABLE_ALIGNMENT.LEFT; intro.autofit = False
for r, (a,b) in enumerate([("Responsável", "[Preencher nome]"), ("Disciplina", "Desenvolvimento 2"), ("Versão", "1.0 | [Preencher data]" )]):
    for c, v in enumerate((a,b)):
        cell = intro.rows[r].cells[c]; set_width(cell, 4.0 if c == 0 else 11.0); set_cell_margins(cell, 130, 160, 130, 160); set_cell_border(cell)
        if c == 0: set_cell_shading(cell, PALE)
        add_text(cell.paragraphs[0], v, bold=(c==0), color=NAVY if c==0 else None, size=10)

doc.add_paragraph().paragraph_format.space_after = Pt(60)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
add_text(p, "Finalidade", bold=True, color=NAVY, size=11)
para(doc, "Este documento registra o escopo, a arquitetura, as evidências e o estado de implementação do CRUD administrativo desenvolvido em Spring Boot. Ele deve ser atualizado à medida que os endpoints e os testes forem concluídos.", space_after=0)

doc.add_page_break()
heading(doc, "1. Contexto e objetivo")
para(doc, "O CRUD administrativo é a camada de back-end responsável por cadastrar, consultar, atualizar e excluir informações administrativas do projeto Mottainai. A solução utiliza Java e Spring Boot para expor uma API REST conectada ao PostgreSQL.")
heading(doc, "Responsabilidade individual", 2)
para(doc, "Responsável pela implementação: [preencher nome]. A responsabilidade inclui estruturar a API, integrar o banco de dados, aplicar validações, documentar os testes e manter as evidências de entrega atualizadas.")
heading(doc, "Decisão técnica", 2)
para(doc, "A implementação foi realizada com Spring Boot, conforme alinhamento com o professor da disciplina. Registre aqui a evidência do acordo: [data, e-mail, mensagem ou ata].")
heading(doc, "Relação com as demandas do projeto")
table(doc, ["Disciplina", "Exigência", "Como esta entrega contribui"], [
    ("Desenvolvimento 2", "API REST em Java, PostgreSQL, JPA, CRUD, MVC, validação e exceções.", "Este é o requisito central atendido pelo CRUD."),
    ("MPTI", "Termos do projeto, ferramenta de planejamento e cronograma.", "Registrar esta tarefa, responsável, prazo, status e links de evidência."),
    ("Entrega geral", "Documento de entregas com links diretos para cada requisito.", "Vincular este documento, o repositório, o PR e os testes."),
], [3.0, 6.3, 6.0])

heading(doc, "2. Tecnologias e configuração")
table(doc, ["Tecnologia", "Uso no projeto"], [
    ("Java 21", "Linguagem da aplicação."),
    ("Spring Boot", "Estrutura da aplicação e execução da API."),
    ("Spring Web MVC", "Controladores e respostas HTTP."),
    ("Spring Data JPA", "Mapeamento e persistência de dados."),
    ("PostgreSQL", "Banco de dados relacional."),
    ("Bean Validation", "Validação dos dados recebidos pela API."),
    ("Gradle", "Gerenciamento de dependências e tarefas de build."),
], [4.0, 11.3])

doc.add_page_break()
heading(doc, "3. Arquitetura da solução")
para(doc, "A aplicação deve manter responsabilidades separadas. O fluxo principal é: requisição HTTP -> controller -> service -> repository -> PostgreSQL. DTOs definem os dados de entrada e saída, evitando expor diretamente as entidades do banco.")
table(doc, ["Pacote", "Responsabilidade", "Evidência esperada"], [
    ("controller", "Receber requisições HTTP e devolver respostas REST.", "Classes @RestController e rotas."),
    ("service", "Aplicar regras de negócio e coordenar operações.", "Classes @Service."),
    ("repository", "Persistir e consultar dados com JPA.", "Interfaces JpaRepository."),
    ("model ou entity", "Mapear tabelas e relacionamentos do banco.", "Classes @Entity."),
    ("dto", "Definir contratos de entrada e saída da API.", "RequestDTO e ResponseDTO."),
    ("exception", "Centralizar respostas de erro da aplicação.", "@RestControllerAdvice."),
    ("validation", "Implementar validações específicas, como CPF.", "Anotações e validadores."),
], [3.2, 6.7, 5.4])
heading(doc, "Estado atual do repositório", 2)
para(doc, "Até a data desta versão, o projeto já possui configuração de Spring Boot, dependências de JPA, JDBC, validação e PostgreSQL, script SQL, DTOs e validação de CPF. Controllers, services, entities e repositories precisam estar versionados e funcionais antes de serem declarados como concluídos nesta documentação.")

heading(doc, "4. Recursos e operações CRUD")
para(doc, "Preencha somente recursos realmente implementados e testados. Os exemplos abaixo podem ser ajustados às entidades escolhidas.")
table(doc, ["Recurso", "Criar", "Listar", "Buscar por ID", "Atualizar", "Excluir"], [
    ("Empresa", "POST /api/companies", "GET /api/companies", "GET /api/companies/{id}", "PUT /api/companies/{id}", "DELETE /api/companies/{id}"),
    ("Funcionário", "[preencher]", "[preencher]", "[preencher]", "[preencher]", "[preencher]"),
    ("Endereço ou Loja", "[preencher]", "[preencher]", "[preencher]", "[preencher]", "[preencher]"),
], [2.0, 3.0, 3.0, 3.3, 3.3, 3.5])

doc.add_page_break()
heading(doc, "5. Validação e tratamento de erros")
heading(doc, "Validações", 2)
bullets(doc, [
    "Campos obrigatórios: informar quais campos não podem ser vazios.",
    "E-mail: validar formato e garantir unicidade quando aplicável.",
    "CPF e CNPJ: validar tamanho, caracteres numéricos e regra específica adotada.",
    "Relacionamentos: verificar se os identificadores de entidades relacionadas existem antes de salvar.",
    "Regras de negócio: documentar limites, datas e demais restrições implementadas.",
])
heading(doc, "Respostas de erro", 2)
table(doc, ["Situação", "HTTP", "Resposta esperada"], [
    ("Dados inválidos", "400", "Mensagem por campo para orientar a correção."),
    ("Registro não encontrado", "404", "Mensagem que identifica o recurso e o ID solicitado."),
    ("Conflito de dados", "409", "Mensagem para e-mail, CPF, CNPJ ou outro valor duplicado."),
    ("Erro inesperado", "500", "Mensagem segura para o cliente e log técnico no servidor."),
], [5.0, 2.0, 8.1])
heading(doc, "Exemplo de evidência de teste", 2)
para(doc, "Inclua uma captura ou link do Postman/Insomnia para um cadastro bem-sucedido, uma consulta e uma resposta de validação. Não inclua senhas, tokens ou credenciais reais.")

heading(doc, "6. Banco de dados")
para(doc, "O banco utilizado é PostgreSQL. O script oficial encontra-se em src/main/resources/script.sql e contém a estrutura das tabelas administrativas, restrições de integridade e carga inicial de dados.")
table(doc, ["Variável", "Finalidade", "Exemplo seguro"], [
    ("DB_URL", "URL de conexão com o PostgreSQL.", "jdbc:postgresql://localhost:5432/mottainai"),
    ("DB_USERNAME", "Usuário do banco.", "[definido no ambiente]"),
    ("DB_PASSWORD", "Senha do banco.", "[definida no ambiente - não versionar]"),
], [3.5, 5.7, 5.9])

doc.add_page_break()
heading(doc, "7. Como executar")
table(doc, ["Passo", "Ação"], [
    ("1", "Criar um banco PostgreSQL local para o projeto."),
    ("2", "Executar o script SQL do projeto para criar as tabelas e a carga inicial."),
    ("3", "Configurar DB_URL, DB_USERNAME e DB_PASSWORD no ambiente local."),
    ("4", "No diretório back-end/mottainai-crud, executar ./gradlew bootRun."),
    ("5", "Testar as rotas com Postman, Insomnia ou Swagger, quando disponível."),
], [2.0, 13.1])
heading(doc, "8. Evidências de entrega")
para(doc, "Use links diretos para que o avaliador consiga verificar cada item sem procurar manualmente no repositório.")
table(doc, ["Evidência", "Link ou local", "Status"], [
    ("Repositório do projeto", "[inserir URL do GitHub]", "[preencher]"),
    ("Pull request ou commits", "[inserir URL]", "[preencher]"),
    ("Script do banco", "src/main/resources/script.sql", "Disponível"),
    ("Documentação da API", "[inserir URL Swagger ou coleção Postman]", "[preencher]"),
    ("Teste de criação", "[inserir evidência]", "[preencher]"),
    ("Teste de atualização e exclusão", "[inserir evidência]", "[preencher]"),
    ("Planejamento MPTI", "[link da tarefa no Jira/Trello/Miro]", "[preencher]"),
], [4.5, 8.0, 2.6])

heading(doc, "9. Resultado e próximos passos")
para(doc, "Resultado atual: a base técnica da aplicação e do banco está estruturada. Antes da entrega final, concluir e versionar os endpoints, a camada de persistência, as regras de negócio, o tratamento centralizado de exceções e os testes de todas as operações CRUD.")
bullets(doc, [
    "Versionar as classes pendentes e manter os commits descritivos.",
    "Atualizar a tabela de endpoints apenas com rotas em funcionamento.",
    "Anexar evidências de testes de sucesso e de erro.",
    "Vincular a tarefa ao cronograma e à ferramenta de planejamento de MPTI.",
    "Atualizar esta documentação antes da apresentação final.",
])

heading(doc, "Referência de demanda", 2)
para(doc, "Projeto Interdisciplinar - Solicitações por Disciplina, 14/05/2026. Requisitos utilizados: Metodologia de Projetos de TI e Desenvolvimento 2.", space_after=0)

doc.core_properties.title = "Documentação CRUD administrativo com Spring Boot"
doc.core_properties.subject = "Projeto Mottainai"
doc.core_properties.author = "Mottainai"
doc.save(OUT)
print(OUT)
